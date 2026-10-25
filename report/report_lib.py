"""Small APA-7 document builder on top of python-docx.

Uses the submitted synopsis as the style base (Times New Roman 12, 1-inch margins, page-number
header), clears its body and writes new content. Inline markup in text:
    *italic*   **bold**   [[placeholder to fill]]  (placeholders are highlighted yellow)
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\[\[.+?\]\])")


class APADoc:
    def __init__(self, base: str):
        self.doc = Document(base)
        body = self.doc.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        self._num_counter = 100
        self._ensure_numbering()
        self.headings: list[tuple[int, str]] = []
        self.placeholders = 0
        self.body_words = 0
        self._counting = True
        # ask Word to refresh fields (TOC) on open
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            uf = OxmlElement("w:updateFields")
            uf.set(qn("w:val"), "true")
            after = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                     "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                     "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
                     "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                     "decimalSymbol", "listSeparator"]
            anchor = next((el for el in settings if el.tag.split("}")[1] in after), None)
            if anchor is not None:
                anchor.addprevious(uf)
            else:
                settings.append(uf)

    # ---------------------------------------------------------------- numbering
    def _ensure_numbering(self):
        numbering = self.doc.part.numbering_part.element
        def abstract(aid, fmt, text):
            a = OxmlElement("w:abstractNum")
            a.set(qn("w:abstractNumId"), str(aid))
            m = OxmlElement("w:multiLevelType"); m.set(qn("w:val"), "singleLevel"); a.append(m)
            lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
            for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", text), ("w:lvlJc", "left")):
                e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
            ppr = OxmlElement("w:pPr"); ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "360"); ppr.append(ind); lvl.append(ppr)
            if fmt == "bullet":
                rpr = OxmlElement("w:rPr"); f = OxmlElement("w:rFonts")
                f.set(qn("w:ascii"), "Symbol"); f.set(qn("w:hAnsi"), "Symbol"); f.set(qn("w:hint"), "default")
                rpr.append(f); lvl.append(rpr)
            a.append(lvl)
            return a
        first_num = numbering.find(qn("w:num"))
        for a in (abstract(90, "bullet", ""), abstract(91, "decimal", "%1.")):
            if first_num is not None:
                first_num.addprevious(a)
            else:
                numbering.append(a)
        self._numbering = numbering

    def _new_num(self, abstract_id: int) -> int:
        self._num_counter += 1
        n = OxmlElement("w:num"); n.set(qn("w:numId"), str(self._num_counter))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), str(abstract_id)); n.append(a)
        o = OxmlElement("w:lvlOverride"); o.set(qn("w:ilvl"), "0")
        s = OxmlElement("w:startOverride"); s.set(qn("w:val"), "1"); o.append(s); n.append(o)
        self._numbering.append(n)
        return self._num_counter

    # ---------------------------------------------------------------- runs
    def _runs(self, par, text: str, size: float | None = None, base_bold=False, base_italic=False):
        for tok in TOKEN_RE.split(text):
            if not tok:
                continue
            bold, italic, ph = base_bold, base_italic, False
            if tok.startswith("**") and tok.endswith("**"):
                tok, bold = tok[2:-2], True
            elif tok.startswith("[[") and tok.endswith("]]"):
                tok, ph = tok[2:-2], True
                self.placeholders += 1
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, italic = tok[1:-1], not base_italic
            # nested italic inside bold e.g. **RQ1 (*Accuracy*)** is not needed; keep simple
            r = par.add_run(tok)
            r.bold = bold or None
            r.italic = italic or None
            if size:
                r.font.size = Pt(size)
            if ph:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.text = f"[{tok}]"
        if self._counting:
            self.body_words += len(re.sub(r"[\[\]*]", "", text).split())
        return par

    @staticmethod
    def _spacing(par, line=480, before=0, after=0):
        pf = par.paragraph_format
        pPr = par._p.get_or_add_pPr()
        sp = pPr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing"); pPr.append(sp)
        sp.set(qn("w:line"), str(line)); sp.set(qn("w:lineRule"), "auto")
        sp.set(qn("w:before"), str(before)); sp.set(qn("w:after"), str(after))
        return pf

    # ---------------------------------------------------------------- blocks
    def p(self, text: str, indent=True, align=None, keep_next=False, size=None, bold=False, italic=False):
        par = self.doc.add_paragraph()
        self._spacing(par)
        if indent:
            par.paragraph_format.first_line_indent = Inches(0.5)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if keep_next:
            par.paragraph_format.keep_with_next = True
        return self._runs(par, text, size=size, base_bold=bold, base_italic=italic)

    def blank(self, n=1):
        for _ in range(n):
            par = self.doc.add_paragraph(); self._spacing(par)

    def page_break(self):
        par = self.doc.add_paragraph()
        par.add_run().add_break(WD_BREAK.PAGE)

    def heading(self, level: int, text: str, toc=True):
        par = self.doc.add_paragraph(style=f"Heading {level}")
        self._spacing(par)
        if level == 1:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        self._runs(par, text)
        if toc and level <= 2:
            self.headings.append((level, text))
        return par

    def h1(self, t, toc=True): return self.heading(1, t, toc)
    def h2(self, t, toc=True): return self.heading(2, t, toc)
    def h3(self, t): return self.heading(3, t, toc=False)

    def run_in(self, lead: str, text: str):
        """APA level-4 style: bold run-in heading ending with a period, then text."""
        return self.p(f"**{lead}** {text}")

    def bullets(self, items, numbered=False):
        num = self._new_num(91 if numbered else 90)
        for it in items:
            par = self.doc.add_paragraph(style="List Paragraph")
            self._spacing(par)
            pPr = par._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); numPr.append(il)
            ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num)); numPr.append(ni)
            ps = pPr.find(qn("w:pStyle"))
            if ps is not None:
                ps.addnext(numPr)
            else:
                pPr.insert(0, numPr)
            self._runs(par, it)

    def label(self, kind: str, num, title: str):
        a = self.doc.add_paragraph(); self._spacing(a, before=240)
        a.paragraph_format.keep_with_next = True
        self._counting = False
        self._runs(a, f"**{kind} {num}**")
        b = self.doc.add_paragraph(); self._spacing(b)
        b.paragraph_format.keep_with_next = True
        self._runs(b, title, base_italic=True)
        self._counting = True

    def note(self, text: str):
        par = self.doc.add_paragraph(); self._spacing(par, line=240, before=60, after=240)
        self._counting = False
        self._runs(par, "*Note.* " + text, size=10)
        self._counting = True

    def table(self, num, title, header, rows, widths, note=None, size=9, group_rows=()):
        """APA table: rules above/below the header and below the last row only."""
        self.label("Table", num, title)
        self._counting = False
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl = t._tbl
        tblPr = tbl.tblPr
        tw = tblPr.find(qn("w:tblW"))
        tw.set(qn("w:w"), str(sum(widths))); tw.set(qn("w:type"), "dxa")
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
        look = tblPr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(lay)
        else:
            tblPr.append(lay)
        grid = tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        all_rows = [header] + rows
        for i, (row, cells) in enumerate(zip(t.rows, all_rows)):
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            if i == 0:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            is_group = i > 0 and (i - 1) in group_rows
            for j, (cell, txt) in enumerate(zip(row.cells, cells)):
                tcPr = cell._tc.get_or_add_tcPr()
                w = tcPr.find(qn("w:tcW"))
                if w is None:
                    w = OxmlElement("w:tcW"); tcPr.insert(0, w)
                w.set(qn("w:w"), str(widths[j])); w.set(qn("w:type"), "dxa")
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    b = OxmlElement(f"w:{side}")
                    rule = (i == 0 and side in ("top", "bottom")) or (i == len(all_rows) - 1 and side == "bottom")
                    b.set(qn("w:val"), "single" if rule else "nil")
                    if rule:
                        b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "000000")
                    borders.append(b)
                tcPr.append(borders)
                if is_group:
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F2"); tcPr.append(shd)
                mar = OxmlElement("w:tcMar")
                for side, v in (("top", 30), ("left", 70), ("bottom", 30), ("right", 70)):
                    m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), str(v)); m.set(qn("w:type"), "dxa"); mar.append(m)
                tcPr.append(mar)
                par = cell.paragraphs[0]
                self._spacing(par, line=240, before=10, after=10)
                if is_group and j > 0:
                    continue
                self._runs(par, str(txt), size=size, base_bold=(i == 0), base_italic=is_group)
            if is_group:
                # merge group row across all columns
                row.cells[0].merge(row.cells[-1])
        self._counting = True
        if note:
            self.note(note)
        else:
            spacer = self.doc.add_paragraph(); self._spacing(spacer, line=240, after=120)

    def figure(self, num, title, image: str | Path | None, placeholder: str = "", note=None, width_in=6.0):
        self.label("Figure", num, title)
        if image and Path(image).exists():
            par = self.doc.add_paragraph(); self._spacing(par, line=240, before=60, after=60)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = bool(note)
            par.add_run().add_picture(str(image), width=Inches(width_in))
        else:
            self._counting = False
            t = self.doc.add_table(rows=1, cols=1)
            cell = t.rows[0].cells[0]
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                b = OxmlElement(f"w:{side}"); b.set(qn("w:val"), "dashed"); b.set(qn("w:sz"), "8"); b.set(qn("w:color"), "7F7F7F")
                borders.append(b)
            tcPr.append(borders)
            par = cell.paragraphs[0]; self._spacing(par, line=240, before=400, after=400)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._runs(par, f"[[{placeholder}]]", size=10)
            self._counting = True
        if note:
            self.note(note)
        else:
            spacer = self.doc.add_paragraph(); self._spacing(spacer, line=240, after=120)

    def reference(self, text: str):
        par = self.doc.add_paragraph(); self._spacing(par)
        par.paragraph_format.left_indent = Inches(0.5)
        par.paragraph_format.first_line_indent = Inches(-0.5)
        self._counting = False
        self._runs(par, text)
        self._counting = True

    def toc(self, entries: list[tuple[int, str, str]] | None):
        """TOC field with cached entries (level, text, page) so it renders before Word updates it."""
        title = self.doc.add_paragraph(); self._spacing(title)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._counting = False
        self._runs(title, "**Table of Contents**")
        entries = entries or [(1, "(Right-click and choose Update Field to build the table of contents.)", "")]
        for k, (lvl, text, page) in enumerate(entries):
            par = self.doc.add_paragraph(style=f"toc {min(lvl, 2)}")
            self._spacing(par, line=240, after=40)
            if k == 0:
                self._fld(par, "begin"); self._instr(par, ' TOC \\o "1-2" \\h \\z \\u '); self._fld(par, "separate")
            par.add_run(re.sub(r"[*\[\]]", "", text))
            par.add_run().add_tab()
            par.add_run(str(page))
            if k == len(entries) - 1:
                self._fld(par, "end")
        self._counting = True

    @staticmethod
    def _fld(par, kind):
        r = par.add_run()
        f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)

    @staticmethod
    def _instr(par, text):
        r = par.add_run()
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = text; r._r.append(i)

    def save(self, path):
        self.doc.save(path)
